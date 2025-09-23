
import requests
import time
from Bio import SeqIO
import os

output_file = r'e:\Guido\sibel\Masters_Thesis\outputs\extracted ref seq.tsv'
fasta_path = r'E:\Guido\sibel\Masters_Thesis\Uniprot_id_matching\search_FASTA_ENB\GCF_904849725.1_Hordeum_vulgare.faa'


# Step 1: Extract XP_... IDs from FASTA
def extract_refseq_ids(fasta_file):
    ids = set()
    for record in SeqIO.parse(fasta_file, "fasta"):
        if record.id.startswith("XP_"):
            ids.add(record.id)
    return list(ids)

# Step 2: Submit mapping job to UniProt
def submit_mapping_job(id_list):
    url = "https://rest.uniprot.org/idmapping/run"
    data = {
        "from": "RefSeq_Protein",
        "to": "UniProtKB",
        "ids": ",".join(id_list)
    }
    response = requests.post(url, data=data)
    response.raise_for_status()
    return response.json()["jobId"]

# Step 3: Wait for job completion
def wait_for_job(job_id):
    url = f"https://rest.uniprot.org/idmapping/status/{job_id}"
    while True:
        print(f"🔍 Checking status for job ID: {job_id}")
        r = requests.get(url)
        
        # If it's not JSON, print raw response
        try:
            data = r.json()
        except Exception as e:
            print("❌ Response is not JSON — dumping raw text:")
            print(r.text)
            raise Exception("Failed to parse JSON from UniProt")

        # Case 1: Expected jobStatus
        if "jobStatus" in data:
            print(f"📦 jobStatus: {data['jobStatus']}")
            if data["jobStatus"] == "FINISHED":
                print("✅ Job finished.")
                break
            elif data["jobStatus"] == "FAILED":
                raise Exception("❌ UniProt job failed.")
            else:
                print("⏳ Job still running...")
                time.sleep(3)

        # Case 2: UniProt skipped jobStatus and gave results directly
        elif "results" in data:
            print("✅ Job finished (results returned directly).")
            break

        # Case 3: Something else entirely
        else:
            print("❗ Unexpected response from UniProt:")
            print(data)  # This will show the actual response dict
            raise Exception("Can't determine job status — check response above")

# Step 4: Download batch result
def download_results(job_id, filename):
    url = f"https://rest.uniprot.org/idmapping/uniprotkb/results/{job_id}?format=tsv"
    r = requests.get(url)
    r.raise_for_status()
    with open(filename, "w") as f:
        f.write(r.text)
    print(f"✅ Saved: {filename}")

# Step 5: Main pipeline
def main():
    output_file = r'e:\Guido\sibel\Masters_Thesis\outputs\extracted ref seq.tsv'
    fasta_path = r'E:\Guido\sibel\Masters_Thesis\Uniprot_id_matching\search_FASTA_ENB\GCF_904849725.1_Hordeum_vulgare.faa'        

    refseq_ids = extract_refseq_ids(fasta_path)
    print(f"Found {len(refseq_ids)} XP_ IDs in {fasta_path}")

    batch_size = 500
    all_results = []

    for i in range(0, len(refseq_ids), batch_size):
        batch = refseq_ids[i:i+batch_size]
        print(f"\n🔁 Batch {i//batch_size + 1}: {len(batch)} IDs")
        try:
            job_id = submit_mapping_job(batch)
            wait_for_job(job_id)
            batch_file = f"batch_{i//batch_size + 1}.tsv"
            download_results(job_id, batch_file)
            all_results.append(batch_file)
        except Exception as e:
            print(f"❌ Error in batch {i//batch_size + 1}:", e)

    # Merge all batch files into final output
    with open(output_file, "w") as outfile:
        header_written = False
        for file in all_results:
            with open(file) as infile:
                for i, line in enumerate(infile):
                    if i == 0 and header_written:
                        continue  # Skip header
                    outfile.write(line)
            header_written = True

    print(f"\n🎉 All batches merged into: {output_file}")

if __name__ == "__main__":
    main()



import pandas as pd
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

# Read your DIAMOND output with headers
input = r'e:\Guido\sibel\matches_Hv (2)_Diamond.csv'
Uniprot_id_matching\search_FASTA_ENB\GCF_904849725.1_Hordeum_vulgare.faa
output = r'e:\Guido\sibel\import\Uniprot ENB Id matching\top_blast_hits_cleaned_Hv1.tsv'

df = pd.read_csv(input, sep='\t')

# Clean up column names (remove spaces, unify casing for convenience)
df.columns = [col.strip().replace(' ', '_').lower() for col in df.columns]
print("read and complete")
print(df.columns.tolist())


# Sort to prioritize top hits: lowest e-value, then highest bit-score
df_sorted = df.sort_values(by=['cseqid', 'e-value', 'bit-score'], ascending=[True, True, False])
print("filtering ...")
# Keep only the top hit per cseqid
df_top = df_sorted.drop_duplicates(subset='cseqid', keep='first')

print("concatinating ...")
# Select columns of interest
result_df = df_top[['cseqid', 'mseqid', 'similarity', 'e-value', 'bit-score', 'mismatches', 'gap_openings']]
print("saving ...")
# Save to new file
result_df.to_csv(output, sep='\t', index=False)

# Show preview
print(result_df.head())

###########################################################################################################################################

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font


# Load the original and new UniProt files
df1_in = r'e:\Guido\sibel\import\Uniprot ENB Id matching\top_blast_hits_cleaned_Hv1.tsv'
df2_in = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\extracted ref seq.tsv'
df1 = pd.read_csv(df1_in, sep='\t')  # original data
df2 = pd.read_csv(df2_in, sep='\t')    # new data from UniProt

# Define the column name that has the UniProt entry numbers
ENTRY_COLUMN = "Entry"  # Replace with actual column name if different

# 🔄 Merge df1 with df2 on 'Entry', keeping all rows from df1
merged_df = df1.merge(df2, on=ENTRY_COLUMN, how='left', suffixes=('', '_df2'))

# Save to Excel (to enable styling)
output_file = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\merged_proteins.xlsx'
merged_df.to_excel(output_file, index=False)

# Reload with openpyxl to apply bold formatting
wb = load_workbook(output_file)
ws = wb.active

# Collect duplicates (i.e., Entry IDs found in both df1 and df2)
duplicate_entries = set(df1[ENTRY_COLUMN]) & set(df2[ENTRY_COLUMN])

# Find the column index for the entry column (1-based for openpyxl)
entry_col_idx = None
for i, cell in enumerate(ws[1], 1):
    if cell.value == ENTRY_COLUMN:
        entry_col_idx = i
        break

# Apply bold formatting to duplicated entry numbers
bold_font = Font(bold=True)
for row in ws.iter_rows(min_row=2):
    cell = row[entry_col_idx - 1]
    if cell.value in duplicate_entries:
        cell.font = bold_font

# Save final styled Excel file
wb.save(r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\merged_refseq with diamond extract.xlsx')

print("✅ Saved formatted Excel file:", wb.path)
# Also save as CSV (no formatting)
merged_df.to_csv(r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\merged_refseq with diamond extract.csv', index=False)

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

# Load the merged Excel file created earlier
merged_file = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\merged_refseq with diamond extract.xlsx'
df = pd.read_excel(merged_file)

# --- Step 2: Define column names ---
FROM_COL = "From"
REFSEQ_COL = "RefSeq"
EXCHANGED_COL = "exchanged"

# --- Step 3: Clean column names (in case of extra spaces) ---
df.columns = df.columns.str.strip()

# --- Step 4: Create new column to store replaced RefSeq values ---
df[EXCHANGED_COL] = ""

# --- Step 5: Compare and swap where needed ---
for idx, row in df.iterrows():
    from_id = str(row.get(FROM_COL, '')).strip()
    refseq_id = str(row.get(REFSEQ_COL, '')).strip()
    
    if from_id and refseq_id and from_id != refseq_id:
        df.at[idx, EXCHANGED_COL] = refseq_id    # Store old RefSeq
        df.at[idx, REFSEQ_COL] = from_id         # Replace with From

# --- Step 6: Handle NaN / empty ---
df[REFSEQ_COL] = df[REFSEQ_COL].replace(["", "nan", "NaN", "None"], pd.NA)

# --- Step 7: If RefSeq is missing, fill with 'exchanged' ---
df.loc[df[REFSEQ_COL].isna(), REFSEQ_COL] = df.loc[df[REFSEQ_COL].isna(), EXCHANGED_COL]

# --- Step 8: Save to TEMP Excel ---
temp_output = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\temp_merged.xlsx'
df.to_excel(temp_output, index=False)

# --- Step 9: Reopen for formatting ---
wb = load_workbook(temp_output)
ws = wb.active

# Get column headers
headers = [cell.value for cell in ws[1]]

# Find column indices (1-based)
refseq_idx = headers.index(REFSEQ_COL) + 1
exchanged_idx = headers.index(EXCHANGED_COL) + 1

# Font: bold + italic
bold_italic = Font(bold=True, italic=True)

# Apply formatting only where a change happened
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    refseq_cell = row[refseq_idx - 1]
    exchanged_cell = row[exchanged_idx - 1]

    # Apply bold+italic if there was a swap
    if exchanged_cell.value not in (None, "", "nan", "NaN"):
        refseq_cell.font = bold_italic

# Save final Excel file
final_output = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\merged_refseq with diamond extract final.xlsx'
wb.save(final_output)
wb.close()



print("✅ Done: Comparison, replacement, and formatting completed.")











































# # matching uniprot entries to refseq on OG combined_proteins data
# import pandas as pd

# # File paths
# ref_seq_map = r'e:\Guido\sibel\export\extract ENB annotation to RefSeq and match to Uniprot ID\extracted ref seq.csv'
# og_data = r'e:\Guido\sibel\export\combined_protein.csv'

# # Load CSVs
# csv1 = pd.read_csv(ref_seq_map)
# csv2 = pd.read_csv(og_data)

# # Clean up whitespace and enforce string type
# csv1['From'] = csv1['From'].astype(str).str.strip()
# csv1['Entry'] = csv1['Entry'].astype(str).str.strip()
# csv1['Entry Name'] = csv1['Entry Name'].astype(str).str.strip()

# csv2['Protein'] = csv2['Protein'].astype(str).str.strip()
# csv2['Protein ID'] = csv2['Protein ID'].astype(str).str.strip()
# csv2['Entry Name'] = csv2['Entry Name'].astype(str).str.strip()

# # Merge on Protein (from csv2) == From (from csv1)
# merged = pd.merge(
#     csv2,
#     csv1[['From', 'Entry', 'Entry Name']],
#     how='left',
#     left_on='Protein',
#     right_on='From'
# )

# # Replace 'Protein ID' with 'Entry' where matched
# merged['Protein ID'] = merged['Entry'].combine_first(merged['Protein ID'])

# # Replace 'Entry Name' with matched name where available
# merged['Entry Name'] = merged['Entry Name_y'].combine_first(merged['Entry Name_x'])

# # Drop temporary columns
# merged.drop(columns=['From', 'Entry', 'Entry Name_x', 'Entry Name_y'], inplace=True)

# num_matches = merged['Protein ID'].isin(csv1['Entry']).sum()
# print(f"🔁 Protein ID updated in {num_matches} rows.")

# # Save result
# output_path = r'e:\Guido\sibel\export\updated_combined_protein.csv'
# merged.to_csv(output_path, index=False)
# print(f"\n✅ File saved to: {output_path}")


